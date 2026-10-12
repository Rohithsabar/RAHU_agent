from __future__ import annotations

import os
import re
import webbrowser
import tempfile
from zipfile import ZIP_DEFLATED, ZipFile
from datetime import datetime
from html import unescape
from pathlib import Path
from urllib.parse import quote

from lxml import etree
import requests

try:
    import pyautogui
except Exception:  # pragma: no cover
    pyautogui = None

try:
    from docx import Document as WordDocument
except Exception:  # pragma: no cover
    WordDocument = None

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR
    from pptx.util import Inches, Pt
except Exception:  # pragma: no cover
    Presentation = None


class DesktopAutomation:
    def __init__(self) -> None:
        self.pyautogui = pyautogui

    def create_word_document(self, title: str, body: str, folder: str = ".") -> str:
        folder_path = Path(folder).expanduser() if folder and folder != "." else Path.cwd()
        folder_path.mkdir(parents=True, exist_ok=True)
        base_name = (title or "document").replace(" ", "_")
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"{base_name}_{timestamp}.docx"
        full_path = folder_path / filename

        if WordDocument is not None:
            doc = WordDocument()
            doc.add_heading(title or "Document", level=1)
            for paragraph in (body or "").split("\n"):
                if paragraph.strip():
                    doc.add_paragraph(paragraph.strip())
            doc.save(str(full_path))
            return str(full_path)

        text_path = folder_path / f"{base_name}_{timestamp}.txt"
        text_path.write_text(body or title or "Document", encoding="utf-8")
        return str(text_path)

    def create_presentation(self, title: str, slides: list[str], folder: str = ".") -> str:
        folder_path = Path(folder).expanduser() if folder and folder != "." else Path.cwd()
        folder_path.mkdir(parents=True, exist_ok=True)
        base_name = (title or "presentation").replace(" ", "_")
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"{base_name}_{timestamp}.pptx"
        full_path = folder_path / filename

        if Presentation is not None:
            prs = Presentation()
            prs.slide_width = Inches(13.333)
            prs.slide_height = Inches(7.5)

            title_slide = prs.slides.add_slide(prs.slide_layouts[0])
            title_slide.background.fill.solid()
            title_slide.background.fill.fore_color.rgb = RGBColor(16, 42, 67)
            title_band = title_slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                0,
                0,
                prs.slide_width,
                Inches(0.18),
            )
            title_band.fill.solid()
            title_band.fill.fore_color.rgb = RGBColor(53, 196, 154)
            title_band.line.fill.background()
            title_slide.shapes.title.text = title or "Presentation"
            title_shape = title_slide.shapes.title
            title_shape.left = Inches(0.9)
            title_shape.top = Inches(1.65)
            title_shape.width = Inches(11.5)
            title_shape.height = Inches(1.5)
            for paragraph in title_shape.text_frame.paragraphs:
                paragraph.font.name = "Aptos Display"
                paragraph.font.size = Pt(34)
                paragraph.font.bold = True
                paragraph.font.color.rgb = RGBColor(255, 255, 255)
            subtitle = title_slide.placeholders[1]
            subtitle.text = "SYSTEM DESIGN  |  WORKFLOW  |  FACILITIES"
            subtitle.left = Inches(0.95)
            subtitle.top = Inches(3.35)
            subtitle.width = Inches(10.8)
            for paragraph in subtitle.text_frame.paragraphs:
                paragraph.font.name = "Aptos"
                paragraph.font.size = Pt(15)
                paragraph.font.color.rgb = RGBColor(155, 223, 205)

            headings = ["Overview", "Key Insights", "Opportunities", "Recommended Actions"]
            for index, item in enumerate(slides or ["Project overview", "Goals", "Next steps"]):
                slide = prs.slides.add_slide(prs.slide_layouts[1])
                requested_heading, separator, requested_body = item.partition("\n")
                default_heading = headings[index] if index < len(headings) else f"Additional Insight {index - len(headings) + 1}"
                heading = requested_heading if separator else default_heading
                body_text = requested_body if separator else item
                points = [point.strip(" -\t") for point in re.split(r"[\n;]+", body_text) if point.strip(" -\t")]
                slide.background.fill.solid()
                slide.background.fill.fore_color.rgb = RGBColor(244, 248, 250) if index % 2 == 0 else RGBColor(235, 245, 241)
                title_bar = slide.shapes.add_shape(
                    MSO_SHAPE.RECTANGLE,
                    0,
                    0,
                    prs.slide_width,
                    Inches(0.12),
                )
                title_bar.fill.solid()
                title_bar.fill.fore_color.rgb = RGBColor(24, 135, 107) if index % 2 == 0 else RGBColor(41, 122, 163)
                title_bar.line.fill.background()

                title_box = slide.shapes.title
                title_box.text = heading
                title_box.left = Inches(0.72)
                title_box.top = Inches(0.38)
                title_box.width = Inches(11.8)
                title_box.height = Inches(0.72)
                for paragraph in title_box.text_frame.paragraphs:
                    paragraph.font.name = "Aptos Display"
                    paragraph.font.size = Pt(25)
                    paragraph.font.bold = True
                    paragraph.font.color.rgb = RGBColor(16, 42, 67)

                if "pipeline" in heading.casefold() and any(">" in point for point in points):
                    pipeline = next(point for point in points if ">" in point)
                    stages = [stage.strip() for stage in pipeline.split(">") if stage.strip()]
                    gap = 0.16
                    box_width = min(1.72, (12.0 - gap * (len(stages) - 1)) / max(len(stages), 1))
                    total_width = box_width * len(stages) + gap * (len(stages) - 1)
                    left = (13.333 - total_width) / 2
                    for stage_index, stage in enumerate(stages):
                        x = Inches(left + stage_index * (box_width + gap))
                        box = slide.shapes.add_shape(
                            MSO_SHAPE.ROUNDED_RECTANGLE,
                            x,
                            Inches(2.35),
                            Inches(box_width),
                            Inches(0.9),
                        )
                        box.fill.solid()
                        box.fill.fore_color.rgb = RGBColor(16, 42, 67) if stage_index % 2 == 0 else RGBColor(24, 135, 107)
                        box.line.fill.background()
                        box.text_frame.text = stage
                        box.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
                        for paragraph in box.text_frame.paragraphs:
                            paragraph.alignment = 1
                            paragraph.font.name = "Aptos"
                            paragraph.font.size = Pt(11)
                            paragraph.font.bold = True
                            paragraph.font.color.rgb = RGBColor(255, 255, 255)
                    remaining_points = [point for point in points if point != pipeline]
                    self._add_slide_bullets(slide, remaining_points, Inches(1.05), Inches(4.0), Inches(11.1))
                else:
                    self._add_slide_bullets(slide, points or [item], Inches(1.05), Inches(1.55), Inches(11.1))

                footer = slide.shapes.add_textbox(Inches(0.75), Inches(7.08), Inches(11.8), Inches(0.25))
                footer.text_frame.text = f"RAHU  /  HOSPITAL SYSTEM DESIGN                                      {index + 2:02d}"
                for paragraph in footer.text_frame.paragraphs:
                    paragraph.font.name = "Aptos"
                    paragraph.font.size = Pt(9)
                    paragraph.font.color.rgb = RGBColor(94, 116, 132)

            prs.save(str(full_path))
            self._add_fade_transitions(full_path)
            return str(full_path)

        text_path = folder_path / f"{base_name}_{timestamp}.txt"
        text_path.write_text("\n".join(slides or ["Project overview", "Goals", "Next steps"]), encoding="utf-8")
        return str(text_path)

    def research_topic(self, topic: str, limit: int = 8) -> list[str]:
        query = (topic or "remote work").strip()
        if not query:
            return []

        url = "https://duckduckgo.com/html/?q=" + quote(query)
        headers = {"User-Agent": "Mozilla/5.0"}
        try:
            response = requests.get(url, timeout=12, headers=headers)
            response.raise_for_status()
            text = response.text
        except Exception:
            return [
                f"Research for {query} should focus on productivity, planning, and measurable results.",
                f"Key areas to explore: goals, timelines, risks, and practical outcomes for {query}.",
                f"Use examples, trends, and recommendations when building the slide deck for {query}.",
            ]

        matches = re.findall(r"result__snippet.*?>(.*?)</a>", text, flags=re.DOTALL | re.IGNORECASE)
        snippets = []
        for snippet in matches:
            cleaned = re.sub(r"<.*?>", " ", snippet)
            cleaned = unescape(cleaned)
            cleaned = " ".join(cleaned.split())
            if cleaned and len(cleaned) > 20:
                snippets.append(cleaned)

        if not snippets:
            return [
                f"A strong overview of {query} includes definition, benefits, challenges, and actionable steps.",
                f"Focus the presentation on outcomes, examples, and useful recommendations for {query}.",
                f"Use real-world evidence and practical conclusions when explaining {query}.",
            ]

        return snippets[:limit]

    def create_presentation_from_topic(
        self,
        topic: str,
        folder: str = ".",
        slide_count: int | None = None,
    ) -> str:
        clean_topic = (topic or "Project Update").strip()
        title = clean_topic.title()
        total_slides = max(3, min(slide_count or 10, 30))
        content_slide_count = total_slides - 1
        if "hospital" in clean_topic.casefold() and any(word in clean_topic.casefold() for word in ("management", "system", "health")):
            sections = [
                ("01  |  Vision and Users", [
                    "Unify patient, clinical, operational, and financial information in one secure system.",
                    "Design around faster care coordination, reliable records, and measurable service quality.",
                    "Patients: registration, appointments, instructions, bills, and results.",
                    "Staff: role-based workflows for clinicians, nurses, reception, pharmacy, diagnostics, finance, and administrators.",
                ]),
                ("02  |  Facilities and Core Modules", [
                    "Front desk: patient identity, registration, scheduling, queue, and admission management.",
                    "Care delivery: outpatient visits, emergency, wards, beds, nursing notes, discharge, and referrals.",
                    "Diagnostics and treatment: laboratory orders/results, imaging, medication, pharmacy, and procedure coordination.",
                    "Operations: billing, insurance, inventory, staff scheduling, reporting, and document management.",
                ]),
                ("03  |  System Design Blueprint", [
                    "Presentation: accessible web and mobile interfaces tailored to each role and facility.",
                    "Application services: identity, scheduling, patient record, clinical orders, inventory, billing, and notifications.",
                    "Data layer: protected patient and operational records, documented retention, backup, and recovery.",
                    "Integration layer: validated interfaces for laboratory, imaging, pharmacy, payment, and external systems.",
                ]),
                ("04  |  Patient Care Pipeline", [
                    "Registration > Triage > Consultation > Orders > Treatment > Billing > Discharge",
                    "Carry the same patient identity and encounter context across every handoff.",
                    "Validate required data and permissions before each step; record changes for audit and continuity.",
                    "Route urgent results and exceptions to the accountable care team, not just a general inbox.",
                ]),
                ("05  |  Design and Delivery Pipeline", [
                    "Discover: interview users, map current journeys, define risks, and document requirements.",
                    "Design: model roles and data; prototype priority screens; test accessibility and workflow fit.",
                    "Build: deliver secure modules in small increments with automated and integration testing.",
                    "Pilot: validate one service area, train users, collect feedback, then phase the rollout.",
                ]),
                ("06  |  Safety, Rollout and Measures", [
                    "Protect health information with encryption, least privilege, audit trails, backup, and incident procedures.",
                    "Use clinical review and approval before enabling decision-support or patient-facing advice.",
                    "Track registration-to-visit time, turnaround time, medication/inventory exceptions, claim rework, uptime, and user feedback.",
                    "Next step: select a pilot department, appoint clinical and technical owners, and baseline current performance.",
                ]),
            ]
            if content_slide_count < len(sections):
                selected = sections[:content_slide_count - 1] + [sections[-1]]
            else:
                selected = sections
            if content_slide_count > len(sections):
                selected.extend(
                    (f"Additional Design Detail {index}", [
                        "Confirm data ownership, workflow exceptions, interface responsibilities, and operational support.",
                        "Review the design with clinical users and update acceptance criteria before implementation.",
                    ])
                    for index in range(1, content_slide_count - len(sections) + 1)
                )
            slides = [f"{heading}\n" + "; ".join(points) for heading, points in selected[:content_slide_count]]
            return self.create_presentation(title, slides, folder=folder)

        research = self.research_topic(clean_topic, limit=8)

        findings = research or [
            f"Define {title} and the main problem or opportunity it addresses.",
            f"Describe the current landscape and why {title} matters to its audience.",
            f"Identify potential benefits and opportunities related to {title}.",
            f"Consider practical challenges, costs, and risks related to {title}.",
        ]
        sections = [
            ("Executive Overview", [f"Purpose: explain {title} and its relevance.", findings[0], "Outcome: leave the audience with clear next steps."]),
            ("Background and Context", [findings[1], f"Define the audience, setting, and need for {title}."]),
            ("Current Landscape", [findings[2] if len(findings) > 2 else f"Review the current state of {title}.", *findings[3:5]]),
            ("Key Evidence", findings[5:8] or [f"Use current, verifiable evidence to assess {title}.", "Validate important claims before making decisions."]),
            ("Opportunities and Benefits", [f"Assess where {title} can create practical value.", "Prioritize opportunities by impact, effort, and audience need."]),
            ("Challenges and Risks", [findings[3] if len(findings) > 3 else f"Identify the main implementation risks for {title}.", "Assign an owner and mitigation to each material risk."]),
            ("Recommended Actions", [f"Set a specific objective for {title}.", "Name an owner, define the first milestone, and confirm required resources."]),
            ("Implementation Roadmap", ["Phase 1: confirm scope, audience, and success criteria.", "Phase 2: deliver a small pilot and collect feedback.", "Phase 3: improve, expand, and review results."]),
            ("Success Measures and Conclusion", ["Track outcomes against a baseline and review them at agreed intervals.", "Decide whether to continue, adjust, or expand based on evidence.", f"Conclusion: move forward with a measurable plan for {title}."]),
        ]
        while len(sections) < content_slide_count:
            sections.insert(-1, (
                f"Additional Consideration {len(sections) - 8}",
                [f"Review audience needs, dependencies, and evidence relevant to {title}.", "Record open questions and assign owners before the next decision."],
            ))
        selected_sections = sections[:content_slide_count]
        if content_slide_count < len(sections) and content_slide_count > 1:
            selected_sections = sections[:content_slide_count - 1] + [sections[-1]]
        slides = [f"{heading}\n" + "; ".join(points) for heading, points in selected_sections]

        return self.create_presentation(title, slides, folder=folder)

    @staticmethod
    def _add_slide_bullets(slide: object, points: list[str], left: object, top: object, width: object) -> None:
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_SHAPE

        for index, point in enumerate(points[:6]):
            y = top + Inches(index * 0.78)
            marker = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, y + Inches(0.1), Inches(0.14), Inches(0.14))
            marker.fill.solid()
            marker.fill.fore_color.rgb = RGBColor(24, 135, 107)
            marker.line.fill.background()
            text_box = slide.shapes.add_textbox(left + Inches(0.28), y, width - Inches(0.3), Inches(0.62))
            text_box.text_frame.word_wrap = True
            text_box.text_frame.text = point
            for paragraph in text_box.text_frame.paragraphs:
                paragraph.font.name = "Aptos"
                paragraph.font.size = Pt(16)
                paragraph.font.color.rgb = RGBColor(39, 61, 77)

    @staticmethod
    def _add_fade_transitions(presentation_path: Path) -> None:
        presentation_ns = "http://schemas.openxmlformats.org/presentationml/2006/main"
        namespace = f"{{{presentation_ns}}}"
        temporary_path = presentation_path.with_suffix(".transition.tmp")
        with ZipFile(presentation_path, "r") as source, ZipFile(temporary_path, "w", ZIP_DEFLATED) as output:
            for item in source.infolist():
                data = source.read(item.filename)
                if re.fullmatch(r"ppt/slides/slide\d+\.xml", item.filename):
                    root = etree.fromstring(data)
                    for old_transition in root.findall(f"{namespace}transition"):
                        root.remove(old_transition)
                    transition = etree.Element(f"{namespace}transition", spd="med", advClick="1")
                    etree.SubElement(transition, f"{namespace}fade")
                    color_map = root.find(f"{namespace}clrMapOvr")
                    insertion_index = root.index(color_map) + 1 if color_map is not None else 1
                    root.insert(insertion_index, transition)
                    data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                output.writestr(item, data)
        os.replace(temporary_path, presentation_path)

    def open_job_search(self, query: str = "remote jobs python developer") -> str:
        keyword = (query or "remote jobs").strip()
        url = "https://www.google.com/search?q=" + quote(keyword)
        webbrowser.open_new_tab(url)
        return f"Opened job search for: {keyword}"

    def type_text(self, text: str) -> str:
        if self.pyautogui is None:
            return "Desktop typing automation is not available in this environment."
        self.pyautogui.write(text, interval=0.05)
        return f"Typed text: {text[:80]}"

    def press_hotkeys(self, *keys: str) -> str:
        if self.pyautogui is None:
            return "Desktop hotkey automation is not available in this environment."
        self.pyautogui.hotkey(*keys)
        return f"Pressed keys: {keys}"

    def take_screenshot(self, file_path: str = "desktop_capture.png") -> str:
        if self.pyautogui is None:
            return "Screenshot automation is not available in this environment."
        path = Path(file_path).expanduser()
        self.pyautogui.screenshot(str(path))
        return f"Saved screenshot to: {path}"
